# -*- coding: utf-8 -*-
from odoo import _
from . import constants as C


def parse_worksheet_strict(ws, target_key):
    """
    Strict parser that enforces exact headers (labels and order).
    Returns (hmap, body) where:
      hmap: {internal_key -> column_index_zero_based}
      body: list of data rows with trimmed strings (header removed)
    """
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append([("" if v is None else str(v).strip()) for v in row])
    if not rows:
        return {}, []

    raw_header = [h if h is not None else "" for h in rows[0]]
    expected_pairs = C.HEADER_LABELS.get(target_key, [])
    expected_labels = [lbl for (_k, lbl) in expected_pairs]

    # Trim trailing empties from header row
    while raw_header and (raw_header[-1] or "").strip() == "":
        raw_header.pop()

    # Column count must match exactly
    if len(raw_header) != len(expected_labels):
        raise ValueError(
            _("Sheet '%(sheet)s' header mismatch.\nExpected exactly %(exp_n)d columns: %(exp)s\n"
              "But found %(got_n)d columns: %(got)s\n\nHeaders cannot be renamed, moved, added or removed.") % {
                "sheet": ws.title,
                "exp_n": len(expected_labels),
                "exp": ", ".join(expected_labels),
                "got_n": len(raw_header),
                "got": ", ".join(raw_header) or "-",
            }
        )

    # Each label must match exactly
    for i, (got, want) in enumerate(zip(raw_header, expected_labels), start=1):
        if (got or "").strip() != want:
            raise ValueError(
                _("Sheet '%(sheet)s' header mismatch at column %(col)d.\n"
                  "Expected: '%(want)s'\nFound:    '%(got)s'\n\nHeaders cannot be renamed, moved, added or removed.") %
                {"sheet": ws.title, "col": i, "want": want, "got": got}
            )

    hmap = {key: idx for idx, (key, _lbl) in enumerate(expected_pairs)}
    body = rows[1:]
    return hmap, body


def build_master_template(env):
    """
    Returns (filename, bytes) of the master import Excel template.
    - Row 1 is read-only (locked) on every sheet (no password prompt).
    - Other cells are editable (unlocked).
    - Dropdown lists that mirror other sheets use IF(TRIM(..)="","",..) so blanks
      don't appear as '0' in validation lists.
    - mrp.bom.finished → Component list = ALL Odoo products + names typed in
      'product.raw' and 'product.semi_finished'.
    - mrp.bom.semi_finished → Product list = ALL semi-finished products
      (ignores purchase_ok/sale_ok/available_in_pos flags) + names typed in its sheet.
    """
    import io
    output = io.BytesIO()

    try:
        import xlsxwriter
        use_xlsxwriter = True
    except Exception:
        use_xlsxwriter = False

    # Models
    Uom = env["uom.uom"]
    PP  = env["product.product"]
    PT  = env["product.template"]
    Cat = env["product.category"]
    Loc = env["stock.location"]

    def _dedup(seq):
        seen = set(); out = []
        for x in seq:
            if x not in seen:
                out.append(x); seen.add(x)
        return out

    # Source data
    uom_names = _dedup(Uom.search([]).mapped("name"))
    existing_categories = _dedup(Cat.search([]).mapped("name"))

    purchasable_variants  = _dedup(PP.search([("purchase_ok", "=", True)]).mapped("name"))
    pos_finished_variants = _dedup(PP.search([("product_tmpl_id.available_in_pos", "=", True)]).mapped("name"))

    # Components for finished BoM: ALL products
    component_variants = _dedup(PP.search([]).mapped("name"))


    if "available_in_pos" in PT._fields:
        semi_domain = [
            "|",
            ("product_type_cm", "=", "semi"),
            "&", "&",
            ("purchase_ok", "=", False),
            ("sale_ok", "=", False),
            ("available_in_pos", "=", False),
        ]
    else:
        semi_domain = [
            "|",
            ("product_type_cm", "=", "semi"),
            "&",
            ("purchase_ok", "=", False),
            ("sale_ok", "=", False),
        ]

    semi_finished_variants = _dedup(PT.search(semi_domain).mapped("name"))

    raw_variants = _dedup(
        PP.search([("product_tmpl_id.product_type_cm", "=", "raw")]).mapped("name")
    )

    # Default semi-finished category (for examples)
    semi_cat_rec = env.ref("import_wizards_caisse_manager_18.category_semi_finished_cm", raise_if_not_found=False)
    semi_cat_name = semi_cat_rec.name if semi_cat_rec else "Produits semi-finis"
    if semi_cat_name and semi_cat_name not in existing_categories:
        existing_categories = [semi_cat_name] + existing_categories

    # Locations
    parent_loc_names = _dedup(Loc.search([]).mapped("name"))
    if "WH" not in parent_loc_names:
        parent_loc_names = ["WH"] + parent_loc_names
    else:
        parent_loc_names = ["WH"] + [n for n in parent_loc_names if n != "WH"]

    MAX_ROWS = 10000

    if use_xlsxwriter:
        import xlsxwriter  # noqa: F401
        wb = xlsxwriter.Workbook(output, {'in_memory': True})

        # Formats
        fmt_header_locked = wb.add_format({
            'bold': True, 'align': 'center', 'valign': 'vcenter',
            'text_wrap': True, 'locked': True,
        })
        fmt_unlocked = wb.add_format({'locked': False})
        fmt_unlocked_center = wb.add_format({'locked': False, 'align': 'center', 'valign': 'vcenter'})
        border_unlocked = wb.add_format({'border': 1, 'locked': False})
        colname = xlsxwriter.utility.xl_col_to_name

        # Helpers
        def _helper_fill_column(ws, col_idx_zero, values):
            """Write helper list in a hidden column for validations."""
            for r, v in enumerate(values, start=1):
                ws.write(r, col_idx_zero, v)
            ws.set_column(col_idx_zero, col_idx_zero, None, None, {'hidden': True})

        def _mirror_or_blank(sheet_name, col_letter, row_1b):
            """
            Mirror 'Sheet'!$A$n but yield "" if blank to avoid '0' in lists.
            """
            return f"=IF(TRIM('{sheet_name}'!${col_letter}${row_1b})=\"\",\"\",'{sheet_name}'!${col_letter}${row_1b})"

        def _protect_sheet_row1_readonly(ws, num_cols, default_col_width=28):
            """
            Row 1 locked, other cells unlocked, no password prompt.
            """
            ws.set_column(0, max(0, max(num_cols - 1, 0) + 20), default_col_width, fmt_unlocked)
            ws.set_row(0, None, fmt_header_locked)
            ws.protect('', {
                'select_locked_cells': False,
                'select_unlocked_cells': True,
                'format_cells': False,
                'format_columns': False,
                'format_rows': False,
                'insert_columns': False,
                'insert_rows': False,
                'delete_columns': False,
                'delete_rows': False,
                'sort': False,
                'autofilter': True,
                'objects': True,
                'scenarios': True,
            })
            ws.freeze_panes(1, 0)

        # Build each sheet
        def _build_sheet(sheet_key):
            title = C.SHEET_TITLES[sheet_key]
            headers = C.HEADER_LABELS.get(sheet_key, [])
            ws = wb.add_worksheet(title)

            # headers (locked)
            for c, (_k, lbl) in enumerate(headers):
                ws.write(0, c, lbl, fmt_header_locked)

            # protect: row 1 locked, others unlocked
            _protect_sheet_row1_readonly(ws, num_cols=len(headers))

            next_hcol = len(headers)
            u0 = uom_names[0] if uom_names else ""

            if sheet_key in ("product.raw", "product.semi_finished"):
                # example row
                if sheet_key == "product.semi_finished":
                    cat_default = semi_cat_name or (existing_categories[0] if existing_categories else "")
                else:
                    cat_default = existing_categories[0] if existing_categories else ""

                example = ["Produit A", u0, "199.99", cat_default, r"mon_image.jpg"]
                for c, v in enumerate(example):
                    ws.write(1, c, v, fmt_unlocked)

                # UoM dropdown (col 1)
                if uom_names:
                    _helper_fill_column(ws, next_hcol, uom_names)
                    ws.data_validation(1, 1, MAX_ROWS, 1, {
                        'validate': 'list',
                        'source': f"${colname(next_hcol)}$1:${colname(next_hcol)}${len(uom_names)}",
                        'ignore_blank': False, 'error_type': 'stop',
                        'error_title': 'Unité invalide',
                        'error_message': "Choisissez l'unité depuis la liste.",
                    })
                next_hcol += 1

                # Categories dropdown (col 3) + mirror from categories sheet
                cat_sheet = C.SHEET_TITLES["product.category"]
                existing_n = len(existing_categories)
                _helper_fill_column(ws, next_hcol, existing_categories)
                for j in range(1, MAX_ROWS + 1):
                    ws.write_formula(existing_n + j, next_hcol, _mirror_or_blank(cat_sheet, "A", j + 1))
                ws.data_validation(1, 3, MAX_ROWS, 3, {
                    'validate': 'list',
                    'source': f"${colname(next_hcol)}$1:${colname(next_hcol)}${existing_n + MAX_ROWS}",
                    'ignore_blank': True, 'error_type': 'stop',
                    'error_title': 'Catégorie invalide',
                    'error_message': "Choisissez une catégorie depuis la liste.",
                })
                next_hcol += 1

                # Image column guidance (col 4)
                img_col = 4
                ws.set_column(img_col, img_col, 36, fmt_unlocked)
                ws.write_comment(0, img_col, (
                    "Image (texte) :\n"
                    "• Saisissez UNIQUEMENT un NOM DE FICHIER (ex: background.png, image.jpg, logo.webp)\n"
                    "• Le fichier doit exister dans le module: static/import_images/"
                ))
                return ws

            if sheet_key == "product.category":
                # example row
                for c, v in enumerate(["Catégorie A", ""]):
                    ws.write(1, c, v, fmt_unlocked)
                existing_n = len(existing_categories)
                _helper_fill_column(ws, next_hcol, existing_categories)
                for j in range(1, MAX_ROWS + 1):
                    ws.write_formula(existing_n + j, next_hcol, _mirror_or_blank(title, "A", j + 1))
                ws.data_validation(1, 1, MAX_ROWS, 1, {
                    'validate': 'list',
                    'source': f"${colname(next_hcol)}$1:${colname(next_hcol)}${existing_n + MAX_ROWS}",
                    'ignore_blank': True, 'error_type': 'stop',
                    'error_title': 'Parent invalide',
                    'error_message': "Choisissez une catégorie parente de la liste.",
                })
                return ws

            if sheet_key == "mrp.bom.semi_finished":
                COMPONENT_ROWS_DEFAULT = 10
                BLOCK_ROWS = 1 + COMPONENT_ROWS_DEFAULT
                START_ROW = 1

                # Product selector must include ALL semi-finished (flags ignored)
                p0 = (semi_finished_variants[0] if semi_finished_variants else "Produit semi-fini A")
                u0 = (uom_names[0] if uom_names else "")
                comp0 = raw_variants[0] if raw_variants else "Matière Première A"
                example = [p0, "BOM-001", "1", u0, "normal", comp0, "2", u0]
                for c, v in enumerate(example):
                    ws.write(START_ROW, c, v, fmt_unlocked)

                # UoM validations (Qty & Component UoM)
                if uom_names:
                    _helper_fill_column(ws, next_hcol, uom_names)
                    ucol = next_hcol
                    for col_idx in (3, 7):
                        ws.data_validation(
                            START_ROW, col_idx, START_ROW + BLOCK_ROWS - 1, col_idx,
                            {
                                'validate': 'list',
                                'source': f"${colname(ucol)}$1:${colname(ucol)}${len(uom_names)}",
                                'ignore_blank': False, 'error_type': 'stop',
                                'error_title': 'Unité invalide',
                                'error_message': "Choisissez l'unité depuis la liste.",
                            }
                        )
                    next_hcol += 1

                # Product (semi-finished) selector (ALL + mirror from its sheet)
                semi_sheet = C.SHEET_TITLES["product.semi_finished"]
                existing_n = len(semi_finished_variants)
                _helper_fill_column(ws, next_hcol, semi_finished_variants)
                for j in range(1, MAX_ROWS + 1):
                    ws.write_formula(existing_n + j, next_hcol, _mirror_or_blank(semi_sheet, "A", j + 1))
                ws.data_validation(START_ROW, 0, START_ROW, 0, {
                    'validate': 'list',
                    'source': f"${colname(next_hcol)}$1:${colname(next_hcol)}${max(1, existing_n + MAX_ROWS)}",
                    'ignore_blank': False, 'error_type': 'stop',
                    'error_title': 'Produit invalide',
                    'error_message': "Choisissez un produit semi-fini existant ou saisi dans l’onglet.",
                })
                next_hcol += 1

                # Components (raw) = purchasable_variants + mirror of product.raw names
                raw_sheet = C.SHEET_TITLES["product.raw"]
                existing_n = len(purchasable_variants)
                _helper_fill_column(ws, next_hcol, purchasable_variants)
                for j in range(1, MAX_ROWS + 1):
                    ws.write_formula(existing_n + j, next_hcol, _mirror_or_blank(raw_sheet, "A", j + 1))
                ws.data_validation(START_ROW, 5, START_ROW + BLOCK_ROWS - 1, 5, {
                    'validate': 'list',
                    'source': f"${colname(next_hcol)}$1:${colname(next_hcol)}${existing_n + MAX_ROWS}",
                    'ignore_blank': False, 'error_type': 'stop',
                    'error_title': 'Composant invalide',
                    'error_message': "Choisissez une MP existante ou saisie dans l’onglet MP.",
                })
                next_hcol += 1

                ws.set_column(4, 4, None, fmt_unlocked, {'hidden': True})
                ws.write(START_ROW, 4, "normal", fmt_unlocked)

                ws.merge_range(START_ROW, 0, START_ROW + BLOCK_ROWS - 1, 0, example[0], fmt_unlocked_center)
                ws.merge_range(START_ROW, 1, START_ROW + BLOCK_ROWS - 1, 1, example[1], fmt_unlocked_center)
                ws.merge_range(START_ROW, 2, START_ROW + BLOCK_ROWS - 1, 2, example[2], fmt_unlocked_center)
                ws.merge_range(START_ROW, 3, START_ROW + BLOCK_ROWS - 1, 3, example[3], fmt_unlocked_center)

                ws.freeze_panes(1, 0)
                for r in range(START_ROW, START_ROW + BLOCK_ROWS):
                    for c in range(0, 8):
                        if r == START_ROW and c < 5:
                            continue
                        ws.write_blank(r, c, None, border_unlocked)
                return ws

            if sheet_key == "mrp.bom.finished":
                COMPONENT_ROWS_DEFAULT = 10
                BLOCK_ROWS = 1 + COMPONENT_ROWS_DEFAULT
                START_ROW = 1

                p0 = (pos_finished_variants[0] if pos_finished_variants else "Produit fini POS A")
                u0 = (uom_names[0] if uom_names else "")
                comp0 = (component_variants[0] if component_variants else "Composant A")
                example = [p0, "BOM-001", "1", u0, "phantom", comp0, "2", u0]
                for c, v in enumerate(example):
                    ws.write(START_ROW, c, v, fmt_unlocked)

                # UoM validations
                if uom_names:
                    _helper_fill_column(ws, next_hcol, uom_names)
                    ucol = next_hcol
                    for col_idx in (3, 7):
                        ws.data_validation(
                            START_ROW, col_idx, START_ROW + BLOCK_ROWS - 1, col_idx,
                            {
                                'validate': 'list',
                                'source': f"${colname(ucol)}$1:${colname(ucol)}${len(uom_names)}",
                                'ignore_blank': False, 'error_type': 'stop',
                                'error_title': 'Unité invalide',
                                'error_message': "Choisissez l'unité depuis la liste.",
                            }
                        )
                    next_hcol += 1

                # Finished product selector (existing PoS finished)
                _helper_fill_column(ws, next_hcol, pos_finished_variants)
                ws.data_validation(START_ROW, 0, START_ROW, 0, {
                    'validate': 'list',
                    'source': f"${colname(next_hcol)}$1:${colname(next_hcol)}${max(1, len(pos_finished_variants))}",
                    'ignore_blank': False, 'error_type': 'stop',
                    'error_title': 'Produit fini invalide',
                    'error_message': "Choisissez un produit fini PoS existant.",
                })
                next_hcol += 1

                # Components list = ALL products + mirrors from product.raw & product.semi_finished
                existing_n = len(component_variants)
                _helper_fill_column(ws, next_hcol, component_variants)

                raw_sheet  = C.SHEET_TITLES["product.raw"]
                semi_sheet = C.SHEET_TITLES["product.semi_finished"]

                for j in range(1, MAX_ROWS + 1):
                    ws.write_formula(existing_n + j,           next_hcol, _mirror_or_blank(raw_sheet,  "A", j + 1))
                    ws.write_formula(existing_n + MAX_ROWS + j, next_hcol, _mirror_or_blank(semi_sheet, "A", j + 1))

                ws.data_validation(START_ROW, 5, START_ROW + BLOCK_ROWS - 1, 5, {
                    'validate': 'list',
                    'source': f"${colname(next_hcol)}$1:${colname(next_hcol)}${max(1, existing_n + 2*MAX_ROWS)}",
                    'ignore_blank': False, 'error_type': 'stop',
                    'error_title': 'Composant invalide',
                    'error_message': "Choisissez un produit existant ou saisi dans les onglets MP/Semi.",
                })
                next_hcol += 1

                ws.set_column(4, 4, None, fmt_unlocked, {'hidden': True})
                ws.write(START_ROW, 4, "phantom", fmt_unlocked)

                ws.merge_range(START_ROW, 0, START_ROW + BLOCK_ROWS - 1, 0, example[0], fmt_unlocked_center)
                ws.merge_range(START_ROW, 1, START_ROW + BLOCK_ROWS - 1, 1, example[1], fmt_unlocked_center)
                ws.merge_range(START_ROW, 2, START_ROW + BLOCK_ROWS - 1, 2, example[2], fmt_unlocked_center)
                ws.merge_range(START_ROW, 3, START_ROW + BLOCK_ROWS - 1, 3, example[3], fmt_unlocked_center)

                ws.freeze_panes(1, 0)
                for r in range(START_ROW, START_ROW + BLOCK_ROWS):
                    for c in range(0, 8):
                        if r == START_ROW and c < 5:
                            continue
                        ws.write_blank(r, c, None, border_unlocked)
                return ws

            if sheet_key == "stock.location":
                for c, v in enumerate(["Test", "WH", "internal"]):
                    ws.write(1, c, v, fmt_unlocked)
                if parent_loc_names:
                    _helper_fill_column(ws, next_hcol, parent_loc_names)
                    ws.data_validation(1, 1, MAX_ROWS, 1, {
                        'validate': 'list',
                        'source': f"${colname(next_hcol)}$1:${colname(next_hcol)}${len(parent_loc_names)}",
                        'ignore_blank': True, 'error_type': 'stop',
                        'error_title': 'Parent invalide',
                        'error_message': 'Choisissez un parent de la liste.',
                    })
                ws.data_validation(1, 2, MAX_ROWS, 2, {
                    'validate': 'list',
                    'source': C.USAGE_CHOICES,
                    'ignore_blank': False, 'error_type': 'stop',
                    'error_title': 'Type invalide',
                    'error_message': 'Choisissez un type de la liste.',
                })
                return ws

        # build all
        for key in C.SHEET_IMPORT_ORDER:
            _build_sheet(key)
        wb.close()

    else:
        # openpyxl fallback (locks row 1; minimal)
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
        from openpyxl.styles import Protection

        wb = Workbook()
        first = True

        def _lock_header_unlock_data(ws, num_cols, max_rows=MAX_ROWS):
            ws.protection.sheet = True
            ws.freeze_panes = "A2"
            unprot = Protection(locked=False)
            for r in range(2, max_rows + 1):
                for c in range(1, num_cols + 1):
                    ws.cell(row=r, column=c).protection = unprot

        for key in C.SHEET_IMPORT_ORDER:
            if first:
                ws = wb.active; ws.title = C.SHEET_TITLES[key]; first = False
            else:
                ws = wb.create_sheet(C.SHEET_TITLES[key])

            headers = C.HEADER_LABELS.get(key, [])
            num_cols = len(headers)

            for c, (_k, lbl) in enumerate(headers, start=1):
                ws.cell(row=1, column=c, value=lbl)
                ws.column_dimensions[get_column_letter(c)].width = 28

            _lock_header_unlock_data(ws, num_cols)

            # minimal examples (unlocked)
            if key == "product.category":
                ex = ["Catégorie A", ""]
            elif key in ("product.raw", "product.semi_finished"):
                u0 = uom_names[0] if uom_names else ""
                if key == "product.semi_finished":
                    c_default = semi_cat_name or (existing_categories[0] if existing_categories else "")
                else:
                    c_default = existing_categories[0] if existing_categories else ""
                ex = ["Produit A", u0, "199.99", c_default, r"mon_image.jpg"]
                ws.column_dimensions[get_column_letter(5)].width = 36
            elif key == "mrp.bom.semi_finished":
                u0 = uom_names[0] if uom_names else ""
                p0 = semi_finished_variants[0] if semi_finished_variants else "Produit Semi-fini A"
                comp0 = raw_variants[0] if raw_variants else "Matière Première A"
                ex = [p0, "BOM-001", "1", u0, "normal", comp0, "2", u0]
            elif key == "mrp.bom.finished":
                u0 = uom_names[0] if uom_names else ""
                comp0 = (component_variants[0] if component_variants else "Composant A")
                ex = ["Produit Fini A", "BOM-001", "1", u0, "phantom", comp0, "2", u0]
            else:
                ex = ["Test", "WH", "internal"]

            for c, v in enumerate(ex, start=1):
                cell = ws.cell(row=2, column=c, value=v)
                cell.protection = Protection(locked=False)

        wb.save(output)

    content = output.getvalue()
    output.close()
    return "Modele_Master_Import.xlsx", content
